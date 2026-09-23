import sys
import os

file_path = r"c:\Users\LENOVO\Desktop\khachsan\S2_final.xlsm"
password = "2804"

print("Checking file:", file_path, "Size:", os.path.getsize(file_path))

# Try standard openpyxl first or msoffcrypto or win32com
try:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, data_only=False)
    print("Sheets in workbook (unencrypted):", wb.sheetnames)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\n--- Sheet: {sheetname} (max_row={ws.max_row}, max_column={ws.max_column}) ---")
        for row in range(1, min(ws.max_row + 1, 30)):
            vals = [ws.cell(row=row, column=col).value for col in range(1, min(ws.max_column + 1, 20))]
            if any(v is not None for v in vals):
                print(f"Row {row}: {vals}")
except Exception as e:
    print("Direct openpyxl failed:", e)

# Try msoffcrypto
try:
    import msoffcrypto
    import io
    print("msoffcrypto is available, attempting decryption with password...")
    with open(file_path, "rb") as f:
        file = msoffcrypto.OfficeFile(f)
        file.load_key(password=password)
        decrypted = io.BytesIO()
        file.decrypt(decrypted)
        decrypted.seek(0)
        import openpyxl
        wb = openpyxl.load_workbook(decrypted, data_only=False)
        print("Decrypted successfully! Sheets:", wb.sheetnames)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            print(f"\n--- Sheet: {sheetname} (max_row={ws.max_row}, max_column={ws.max_column}) ---")
            for row in range(1, min(ws.max_row + 1, 40)):
                vals = [ws.cell(row=row, column=col).value for col in range(1, min(ws.max_column + 1, 20))]
                if any(v is not None for v in vals):
                    print(f"Row {row}: {vals}")
except Exception as e:
    print("msoffcrypto attempt:", e)

# Try win32com
try:
    import win32com.client
    print("Trying win32com...")
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb = excel.Workbooks.Open(file_path, False, True, None, password)
    print("win32com opened workbook successfully!")
    for i in range(1, wb.Sheets.Count + 1):
        ws = wb.Sheets(i)
        print(f"Sheet {i}: {ws.Name}")
    wb.Close(False)
    excel.Quit()
except Exception as e:
    print("win32com attempt:", e)
