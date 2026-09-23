import msoffcrypto
import openpyxl
import os
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"c:\Users\LENOVO\Desktop\khachsan\S2_final.xlsm"
decrypted_path = r"c:\Users\LENOVO\Desktop\khachsan\decrypted_S2.xlsm"
password = "2804"

print(f"Decrypting {file_path} with password '{password}'...")
with open(file_path, "rb") as f_in:
    office_file = msoffcrypto.OfficeFile(f_in)
    office_file.load_key(password=password)
    with open(decrypted_path, "wb") as f_out:
        office_file.decrypt(f_out)

print(f"Decrypted successfully to {decrypted_path} (size: {os.path.getsize(decrypted_path)} bytes)")

# Load with openpyxl
wb = openpyxl.load_workbook(decrypted_path, data_only=False, keep_vba=True)
wb_data = openpyxl.load_workbook(decrypted_path, data_only=True, keep_vba=True)

print(f"Sheet names: {wb.sheetnames}")

dump = {"sheets": []}

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    ws_data = wb_data[sheetname]
    
    print(f"\n==========================================")
    print(f"SHEET: {sheetname} | max_row={ws.max_row}, max_column={ws.max_column}")
    print(f"==========================================")
    
    sheet_dump = {
        "sheetname": sheetname,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "rows": []
    }
    
    for r in range(1, min(ws.max_row + 1, 100)):
        row_vals = []
        row_formulas = []
        has_val = False
        for c in range(1, min(ws.max_column + 1, 30)):
            cell_formula = ws.cell(row=r, column=c).value
            cell_val = ws_data.cell(row=r, column=c).value
            
            if cell_val is not None or cell_formula is not None:
                has_val = True
                
            row_vals.append(str(cell_val) if cell_val is not None else "")
            row_formulas.append(str(cell_formula) if cell_formula is not None else "")
            
        if has_val:
            print(f"Row {r:02d}: {row_vals[:16]}")
            # print formulas if different
            formulas_only = [f for f in row_formulas if str(f).startswith("=")]
            if formulas_only:
                print(f"   Formulas in row {r}: {formulas_only}")
            sheet_dump["rows"].append({
                "row": r,
                "values": row_vals,
                "formulas": row_formulas
            })
            
    dump["sheets"].append(sheet_dump)

with open(r"c:\Users\LENOVO\Desktop\khachsan\s2_full_dump.json", "w", encoding="utf-8") as f_json:
    json.dump(dump, f_json, ensure_ascii=False, indent=2)

print("\nDump complete and saved to s2_full_dump.json")
